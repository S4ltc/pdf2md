"""Gemeinsame Hilfen fuer die Tests: Arbeitsordner in tmp_path und selbst erzeugte Testdateien.

Die Tests brauchen keine echten Buecher. PDFs werden mit fpdf2 erzeugt (mit der mitgelieferten DejaVu-Schrift, damit
Umlaute und Sonderzeichen funktionieren), DOCX mit python-docx, EPUB von Hand."""

import sys
import zipfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest

WURZEL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(WURZEL))

import pdf2md  # noqa: E402

DEJAVU = WURZEL / "schriften" / "DejaVuSans.ttf"
FUELLTEXT = ("Dies ist ein Absatz mit genug Text, damit die Erkennung von Scans nicht anschlaegt und das Buch als "
             "Buch mit Text gilt. ") * 6


@pytest.fixture(autouse=True)
def keine_echte_ablage(tmp_path_factory, monkeypatch):
    """Sicherheitsnetz: kein Test schreibt in die echte Ablage neben pdf2md.py. Frueher lagen danach Testreste
    ("Prüfen/quelle (2).pdf", Protokoll.csv) im Projektordner. Tests mit eigener Ablage (arbeitsordner, test_ui_app)
    setzen die Pfade danach selbst."""
    basis = tmp_path_factory.mktemp("ablage")
    for name, unter in (("EINGANG", "Eingang"), ("FERTIG", "Fertig"), ("PRUEFEN", "Prüfen"),
                        ("PROTOKOLL", "Protokoll.csv"), ("SICHERUNG", "Sicherung")):
        monkeypatch.setattr(pdf2md, name, basis / unter)


@pytest.fixture
def arbeitsordner(tmp_path, monkeypatch):
    """Leitet Eingang, Fertig, Pruefen, Protokoll und Sicherung in einen temporaeren Ordner um und schaltet alles
    ab, was Netz oder viel Zeit braucht (Crossref, Formelreparatur)."""
    ordner = SimpleNamespace(basis=tmp_path, eingang=tmp_path / "Eingang", fertig=tmp_path / "Fertig",
                             pruefen=tmp_path / "Prüfen", protokoll=tmp_path / "Protokoll.csv",
                             sicherung=tmp_path / "Sicherung")
    ordner.eingang.mkdir()
    for name, wert in (("EINGANG", ordner.eingang), ("FERTIG", ordner.fertig), ("PRUEFEN", ordner.pruefen),
                       ("PROTOKOLL", ordner.protokoll), ("SICHERUNG", ordner.sicherung)):
        monkeypatch.setattr(pdf2md, name, wert)
    monkeypatch.setattr(pdf2md, "ONLINE_ABGLEICH", False)
    monkeypatch.setattr(pdf2md, "FORMELN_REPARIEREN", False)
    monkeypatch.setattr(pdf2md, "PDF_ENGINE", "pdfium")
    monkeypatch.setattr(pdf2md, "UEBERSCHRIFTEN_AUS_LESEZEICHEN", True)
    return ordner


@pytest.fixture(scope="session")
def konverter():
    return pdf2md.MarkItDown(enable_plugins=False)


@pytest.fixture
def lauf():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def erzeuge_pdf(pfad, seiten, titel=None, autor=None, lesezeichen=None):
    """Schreibt eine PDF. seiten: Liste von Seitentexten (Zeilenumbruch = neuer Absatz, "" = leere Seite wie bei
    einem Scan). lesezeichen: {Seitenindex: [(Ebene, Name), ...]}."""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_font("DejaVu", fname=str(DEJAVU))
    pdf.set_font("DejaVu", size=11)
    if titel is not None:
        pdf.set_title(titel)
    if autor is not None:
        pdf.set_author(autor)
    for i, text in enumerate(seiten):
        pdf.add_page()
        for ebene, name in (lesezeichen or {}).get(i, []):
            pdf.start_section(name, level=ebene)
        for absatz in text.split("\n") if text else []:
            pdf.multi_cell(0, 6, absatz, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(pfad))
    return pfad


def erzeuge_zweispaltig(pfad, links, rechts, rechts_zuerst=False):
    """Eine Seite mit zwei Textspalten. In welcher Reihenfolge das PDF die Spalten zeichnet (und PDFium sie
    deshalb liefert), bestimmt rechts_zuerst; gelesen wird immer links vor rechts."""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_font("DejaVu", fname=str(DEJAVU))
    pdf.set_font("DejaVu", size=11)
    pdf.add_page()

    def spalte(zeilen, x):
        for k, zeile in enumerate(zeilen):
            pdf.set_xy(x, 30 + 6 * k)
            pdf.cell(85, 6, zeile)
    reihenfolge = [(rechts, 110), (links, 15)] if rechts_zuerst else [(links, 15), (rechts, 110)]
    for zeilen, x in reihenfolge:
        spalte(zeilen, x)
    pdf.output(str(pfad))
    return pfad


def erzeuge_gestaltet(pfad, seiten, lesezeichen=None):
    """PDF mit Schriftgroessen und Fett/Kursiv (Helvetica, eine der Standardschriften jedes PDF-Programms).
    seiten: je Seite eine Liste von (Text, Groesse); im Text markieren **...** fett und __...__ kursiv (fpdf2)."""
    from fpdf import FPDF
    pdf = FPDF()
    for i, zeilen in enumerate(seiten):
        pdf.add_page()
        for ebene, name in (lesezeichen or {}).get(i, []):
            pdf.start_section(name, level=ebene)
        for text, groesse in zeilen:
            pdf.set_font("Helvetica", size=groesse)
            pdf.multi_cell(0, groesse * 0.5, text, markdown=True, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(pfad))
    return pfad


def spaltenzeilen(name, anzahl=20):
    return [f"{name} Spalte, Zeile {k:02d}: der Text" for k in range(1, anzahl + 1)]


def erzeuge_buch(pfad, titel="Testbuch", autor="Anna Beispiel", jahr="2019"):
    """Ein vollstaendiges kleines 'Buch': Metadaten, Copyright-Vermerk, genug Text."""
    return erzeuge_pdf(pfad, [f"Titelseite\n© {jahr} Beispielverlag", FUELLTEXT], titel=titel, autor=autor)


def erzeuge_docx(pfad, titel=None, autor=None, text=FUELLTEXT):
    from docx import Document
    dokument = Document()
    dokument.core_properties.title = titel or ""
    dokument.core_properties.author = autor or ""
    dokument.add_paragraph(text)
    dokument.save(str(pfad))
    return pfad


def erzeuge_epub(pfad, titel, autor, datum, isbn=None, text=FUELLTEXT):
    with zipfile.ZipFile(pfad, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                   '</rootfiles></container>')
        kennung = f'<dc:identifier id="id">urn:isbn:{isbn}</dc:identifier>' if isbn else '<dc:identifier id="id">x</dc:identifier>'
        z.writestr("OEBPS/content.opf",
                   '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
                   f'<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>{titel}</dc:title>'
                   f'<dc:creator>{autor}</dc:creator><dc:date>{datum}</dc:date>{kennung}</metadata>'
                   '<manifest><item id="k1" href="k1.xhtml" media-type="application/xhtml+xml"/></manifest>'
                   '<spine><itemref idref="k1"/></spine></package>')
        z.writestr("OEBPS/k1.xhtml", '<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><head>'
                   f'<title>K1</title></head><body><h1>Kapitel 1</h1><p>{text}</p></body></html>')
    return pfad
