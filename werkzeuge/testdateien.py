"""Erfundene Testdateien fuer den Funktionstest des gebauten Programms: werkzeuge/bauen.py (starttest) legt sie in den
Ordner testdateien neben die Ergebnisdatei, ui_app._starttest wandelt sie in Phase 2 ueber die Api um.

Eine PDF (Metadaten, ©-Vermerk, mehrere Seiten mit Lesezeichen, eine kleine Gittertabelle) und eine DOCX
(Kerneigenschaften, ©-Vermerk). Beide gehoeren nach Fertig: Titel, Autor und Jahr sind vorhanden, eine ISBN oder DOI
gibt es nicht, damit kein Netzabruf noetig ist. Braucht fpdf2 und python-docx (requirements-dev.txt), die nur in der
Build-Umgebung liegen, nicht im Programm. Vorbild: die Erzeuger in tests/conftest.py.

Aufruf (Werkzeug):  python werkzeuge/testdateien.py ORDNER"""
import sys
from pathlib import Path

WURZEL = Path(__file__).resolve().parent.parent
DEJAVU = WURZEL / "schriften" / "DejaVuSans.ttf"     # eingebettete Schrift wie in echten Buechern (Umlaute, ©, ³)

AUTORIN = "Erika Muster"
# Je Testdatei: Titel, Autor und Jahr, wie sie im Kopfblock stehen sollen, und Felder, die die Umwandlung dort
# eintraegt (bauen.funktion_bewerten). Verschiedene Namen, damit Rueckgaengig beide .md neben ihr Original legt.
DATEIEN = {
    "beispielbuch.pdf": {"titel": "Beispielbuch", "autor": AUTORIN, "jahr": "2019",
                         "felder": ("lesezeichen", "tabellen")},
    "beispieldokument.docx": {"titel": "Beispieldokument", "autor": AUTORIN, "jahr": "2020", "felder": ()},
}
ABSATZ = ("Dieser Absatz ist frei erfunden und dient nur dem Funktionstest des gebauten Programms. Er ist lang genug, "
          "damit die Erkennung von Scans nicht anschlaegt und die Datei als Dokument mit Text gilt. ") * 3
KAPITEL = ("Einleitung", "Grundlagen", "Beispiele")
TABELLE = (("Werkstoff", "Dichte", "Einheit"), ("Beispielstahl", "7,85", "g/cm³"), ("Beispielmetall", "2,70", "g/cm³"))


def pdf_anlegen(pfad: Path) -> Path:
    """Titelseite mit ©-Vermerk, danach je Kapitel eine Seite mit Lesezeichen, Kolumnentitel und Seitenzahl in der
    Fusszeile; im zweiten Kapitel eine Tabelle aus umrandeten Zellen (fpdf2 zeichnet Rechtecke, tabellen.py erkennt
    daraus ein Gitter)."""
    from fpdf import FPDF
    angaben = DATEIEN["beispielbuch.pdf"]

    class Buch(FPDF):
        def header(self):
            if self.page_no() > 1:                       # Kolumnentitel: wiederholt sich, faellt als Kopfzeile weg
                self.set_font("DejaVu", size=8)
                self.cell(0, 6, f"{angaben['autor']}: {angaben['titel']}", align="R", new_x="LMARGIN", new_y="NEXT")
                self.ln(4)

        def footer(self):
            if self.page_no() > 1:                       # gedruckte Seitenzahl = PDF-Seite
                self.set_y(-15)
                self.set_font("DejaVu", size=8)
                self.cell(0, 6, str(self.page_no()), align="C")

    pdf = Buch()
    pdf.add_font("DejaVu", fname=str(DEJAVU))
    pdf.set_font("DejaVu", size=11)
    pdf.set_title(angaben["titel"])
    pdf.set_author(angaben["autor"])
    pdf.add_page()
    pdf.set_font("DejaVu", size=20)
    pdf.cell(0, 12, angaben["titel"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", size=11)
    pdf.cell(0, 8, angaben["autor"], new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, f"© {angaben['jahr']} Beispielverlag, (c) {angaben['jahr']} {angaben['autor']}",
             new_x="LMARGIN", new_y="NEXT")
    for nummer, kapitel in enumerate(KAPITEL, 1):
        pdf.add_page()
        pdf.start_section(kapitel, level=0)
        pdf.set_font("DejaVu", size=14)
        pdf.cell(0, 10, kapitel, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", size=11)
        pdf.multi_cell(0, 6, f"Im Kapitel {kapitel} steht nur Fülltext. " + ABSATZ, new_x="LMARGIN", new_y="NEXT")
        if nummer == 2:
            pdf.ln(4)
            pdf.cell(0, 8, "Tabelle 1 – Erfundene Kennwerte", new_x="LMARGIN", new_y="NEXT")
            for zeile in TABELLE:
                for text in zeile:
                    pdf.cell(45, 8, text, border=1)
                pdf.ln(8)
            pdf.ln(4)
            pdf.multi_cell(0, 6, "Text nach der Tabelle.", new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(pfad))
    return pfad


def docx_anlegen(pfad: Path) -> Path:
    from docx import Document
    angaben = DATEIEN["beispieldokument.docx"]
    dokument = Document()
    dokument.core_properties.title = angaben["titel"]
    dokument.core_properties.author = angaben["autor"]
    dokument.add_heading(angaben["titel"], level=1)
    dokument.add_paragraph(f"© {angaben['jahr']} {angaben['autor']}")
    for kapitel in KAPITEL:
        dokument.add_heading(kapitel, level=2)
        dokument.add_paragraph(f"Im Kapitel {kapitel} steht nur Fülltext. " + ABSATZ)
    dokument.save(str(pfad))
    return pfad


def anlegen(ordner: Path) -> list[Path]:
    """Legt beide Testdateien im Ordner an (er wird bei Bedarf erzeugt) und gibt ihre Pfade zurueck."""
    ordner = Path(ordner)
    ordner.mkdir(parents=True, exist_ok=True)
    return [pdf_anlegen(ordner / "beispielbuch.pdf"), docx_anlegen(ordner / "beispieldokument.docx")]


if __name__ == "__main__":
    for datei in anlegen(Path(sys.argv[1] if len(sys.argv) > 1 else "testdateien")):
        print(datei)
