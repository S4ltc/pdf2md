"""Haertetest: erzeugt Testdateien aller unterstuetzten Formate mit Grenzfaellen und laesst sie wie im Fenster
durchlaufen (ui_app.Api: Hinzufuegen, Starten). Am Ende je Datei Ergebnis, Dauer und Kopfblock.

Grenzfaelle: leer, kaputt, passwortgeschuetzt, nur Bild, 2000 Seiten, 60.000 Tabellenzeilen, 30.000 HTML-Absaetze,
Windows-1252 ohne Angabe, Platzhalter-Metadaten ("Microsoft Word - Dokument1", Autor "admin"), Titel nur aus
Sonderzeichen, Unicode- und ueberlange Dateinamen, nicht unterstuetzte Formate. Alles in <ziel>, nichts ausserhalb;
dauert gut eine Minute. Nur erzeugte Dateien (kein Buchtext), das Ergebnis darf trotzdem nicht ins Repo.

Aufruf:  python werkzeuge/haertetest.py <ziel>     (zum Beispiel ein Ordner im Scratchpad, nie dist)"""
import io
import json
import shutil
import sys
import time
import zipfile
from pathlib import Path

WURZEL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(WURZEL))
sys.path.insert(0, str(WURZEL / "tests"))

LANG = ("Dies ist ein längerer Absatz über Werkstoffkunde, Festigkeit und Bruchmechanik mit Umlauten (äöüß) "
        "und Sonderzeichen (± ≤ ≥ µ Ω °C), damit der Text realistisch wirkt. ") * 4


def korpus_erzeugen(ziel: Path) -> None:
    from docx import Document
    from docx.shared import Pt
    from fpdf import FPDF
    from openpyxl import Workbook
    from PIL import Image
    from pptx import Presentation
    from pypdf import PdfReader, PdfWriter

    from conftest import erzeuge_epub, erzeuge_pdf

    def pdf(name, seiten, **kw):
        return erzeuge_pdf(ziel / name, seiten, **kw)

    buch = pdf("PDF normales Buch.pdf", ["Titelseite\n© 2021 Beispielverlag\nISBN 978-3-16-148410-0", LANG, LANG],
               titel="Beispielwerkstoffe kompakt", autor="Lena Probst")
    pdf("PDF ohne Metadaten.pdf", ["Ein Skript ohne Angaben", LANG])
    pdf("PDF Titel nur Sonderzeichen.pdf", ["© 2020 Verlag", LANG], titel='???:*<>|"', autor="Max Muster")
    pdf("PDF Titel mit Pfadzeichen.pdf", ["© 2019 Verlag", LANG], titel='Teil 1/2: A\\B "C" <D>?', autor="Eva Weiß")
    pdf("PDF Platzhalter.pdf", ["© 2018 Verlag", LANG], titel="Microsoft Word - Dokument1", autor="admin")
    pdf("PDF sehr langer Titel.pdf", ["© 2022 Verlag", LANG],
        titel="Ein außerordentlich langer Buchtitel über Konstruktion, Berechnung und Gestaltung von "
              "Maschinenelementen – mit Übungen, Lösungen und vielen Beispielen aus der Praxis für Studium und Beruf",
        autor="Anna Beispiel, Bernd Zweiter, Carla Dritte")
    pdf("PDF leer (Scan).pdf", ["", "", ""], titel="Laborbericht", autor="Gruppe 4")
    pdf("Übung – Ä Ö ü ß 日本語 🙂.pdf", ["© 2017 Verlag", LANG], titel="Unicode Test", autor="Jörg Ärger")
    pdf("PDF " + "sehr langer Dateiname " * 8 + ".pdf", ["© 2016 Verlag", LANG], titel="Langer Name", autor="L. Ang")
    pdf("PDF 2000 Seiten.pdf", [f"Seite {i}: {LANG[:400]}" for i in range(2000)], titel="Dicker Wälzer",
        autor="Viel Schreiber")
    bild = io.BytesIO()
    Image.new("RGB", (800, 1100), "white").save(bild, format="PNG")
    d = FPDF()
    d.set_title("Gescanntes Skript")
    d.set_author("Prof. Scan")
    for _ in range(3):
        d.add_page()
        d.image(io.BytesIO(bild.getvalue()), x=0, y=0, w=210)
    d.output(str(ziel / "PDF nur Bild.pdf"))
    w = PdfWriter()
    for s in PdfReader(buch).pages:
        s.rotate(90)
        w.add_page(s)
    w.add_metadata({"/Title": "Gedrehtes Buch", "/Author": "Rita Dreher"})
    w.write(ziel / "PDF gedrehte Seiten.pdf")
    for name, benutzer in (("PDF Passwort.pdf", "geheim"), ("PDF nur Besitzerpasswort.pdf", "")):
        w = PdfWriter(clone_from=buch)
        w.add_metadata({"/Title": "Geschütztes Buch", "/Author": "Sven Sicher"})
        w.encrypt(user_password=benutzer, owner_password="besitzer", algorithm="AES-256")
        w.write(ziel / name)
    (ziel / "PDF abgeschnitten.pdf").write_bytes(buch.read_bytes()[: buch.stat().st_size // 2])
    (ziel / "PDF null Bytes.pdf").write_bytes(b"")
    (ziel / "PDF ist eigentlich HTML.pdf").write_text("<html><body>Kein PDF</body></html>", encoding="utf-8")

    doc = Document()
    doc.core_properties.title, doc.core_properties.author = "Projektbericht Getriebe", "Tom Bericht"
    doc.add_heading("1 Einleitung", 1)
    doc.add_paragraph(LANG)
    doc.add_paragraph("© 2023 Hochschule Musterstadt")
    tabelle = doc.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            tabelle.cell(r, c).text = f"Z{r}S{c}"
    doc.add_picture(io.BytesIO(bild.getvalue()), width=Pt(100))
    doc.save(ziel / "DOCX normal.docx")
    Document().save(ziel / "DOCX leer.docx")
    gross = Document()
    gross.core_properties.title, gross.core_properties.author = "Sehr langes Dokument", "Viel Schreiber"
    for i in range(6000):
        gross.add_paragraph(f"Absatz {i}: {LANG[:300]}")
    gross.save(ziel / "DOCX 6000 Absaetze.docx")
    (ziel / "DOCX kaputt.docx").write_bytes((ziel / "DOCX normal.docx").read_bytes()[:3000])

    p = Presentation()
    p.core_properties.title, p.core_properties.author = "Vortrag Klebverfahren", "Pia Vortrag"
    folie = p.slides.add_slide(p.slide_layouts[0])
    folie.shapes.title.text = "Klebverfahren im Überblick"
    folie.placeholders[1].text = "© 2024 Lehrstuhl Fügetechnik"
    folie.notes_slide.notes_text_frame.text = "Sprechernotiz äöü"
    p.save(ziel / "PPTX normal.pptx")
    Presentation().save(ziel / "PPTX leer.pptx")

    wb = Workbook()
    wb.active.append(["Probe", "Rm in MPa", "Bemerkung"])
    for i in range(30):
        wb.active.append([f"P{i}", 400 + i, "ok äöü" if i % 3 else None])
    wb.properties.title, wb.properties.creator = "Zugversuche", "Lab Or"
    wb.save(ziel / "XLSX normal.xlsx")
    viel = Workbook()
    for i in range(60000):
        viel.active.append([i, i * 1.5, f"Text {i}", "äöü"])
    viel.save(ziel / "XLSX 60000 Zeilen.xlsx")

    erzeuge_epub(ziel / "EPUB normal.epub", "Grundlagen der Beispielkunde", "Hans Wärme", "2015-03-01", text=LANG)
    with zipfile.ZipFile(ziel / "EPUB kaputt.epub", "w") as z:
        z.writestr("mimetype", "application/epub+zip")

    (ziel / "HTML normal.html").write_text(
        '<html><head><meta charset="utf-8"><title>Leitfaden</title><meta name="author" content="Nora Norm">'
        f'</head><body><p>© 2021 Normverlag</p><p>{LANG}</p></body></html>', encoding="utf-8")
    (ziel / "HTML citation.html").write_text(
        '<html><head><meta name="citation_title" content="Wälzlager im Vergleich">'
        '<meta name="citation_author" content="Muster, Max"><meta name="citation_publication_date" content="2019">'
        f'</head><body><p>{LANG}</p></body></html>', encoding="utf-8")
    (ziel / "HTM windows-1252.htm").write_bytes(
        f"<html><head><title>Größen</title></head><body><p>Übergröße {LANG}</p></body></html>"
        .encode("cp1252", errors="replace"))
    (ziel / "HTML riesig.html").write_text("<html><head><title>Riesig</title></head><body>" + "".join(
        f"<p>Absatz {i} {LANG[:200]}</p>" for i in range(30000)) + "</body></html>", encoding="utf-8")
    (ziel / "nicht unterstuetzt.txt").write_text("Text", encoding="utf-8")
    (ziel / "Altes Word.doc").write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 500)


def durchlaufen(korpus: Path, basis: Path) -> None:
    import pdf2md
    import ui_app
    basis.mkdir()
    (basis / "Einstellungen.json").write_text(json.dumps({"pdf2md.ONLINE_ABGLEICH": False}), encoding="utf-8")
    api = ui_app.Api(basis)
    antwort = api.hinzufuegen([str(p) for p in sorted(korpus.iterdir())])
    print(f"hinzugefügt: {len(antwort['kopiert'])}, abgelehnt: {[n for n, _ in antwort['abgelehnt']]}")
    start = time.time()
    api.starten()
    beginn, stati, nach, ende = {}, {}, 0, None
    while ende is None:
        for e in api.ereignisse(nach)["ereignisse"]:
            nach = e["nr"]
            if e["art"] == "datei_beginnt":
                beginn[e["name"]] = time.time()
            elif e["art"] == "datei_fertig":
                stati[e["name"]] = (time.time() - beginn.get(e["name"], time.time()), e["status"])
            elif e["art"] == "lauf_ende":
                ende = e
        time.sleep(0.1)
    print(f"Lauf: {time.time() - start:.0f} s, {ende['ergebnis']}, Fehler im Lauf: {ende['fehler']}\n")
    for name in sorted(stati):
        dauer, status = stati[name]
        print(f"{dauer:6.1f} s  {status[:150]}")
    print("\nKopfblöcke (Titel | Autor | Jahr | prüfen):")
    for ordner in ("Fertig", "Prüfen"):
        for md in sorted((basis / ordner).glob("*.md")):
            kopf, rumpf = pdf2md.kopf_lesen(md.read_text(encoding="utf-8"))
            print(f"  [{ordner}] {md.stem[:50]:50} | {kopf.get('titel')} | {kopf.get('autor')} | {kopf.get('jahr')}"
                  f" | {kopf.get('pruefen') or ''}"[:220])
    print("\nim Eingang geblieben:", sorted(d.name for d in (basis / "Eingang").iterdir()))


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ziel = Path(sys.argv[1]).resolve()
    if ziel.exists():
        shutil.rmtree(ziel)
    korpus = ziel / "korpus"
    korpus.mkdir(parents=True)
    korpus_erzeugen(korpus)
    durchlaufen(korpus, ziel / "ablage")


if __name__ == "__main__":
    main()
